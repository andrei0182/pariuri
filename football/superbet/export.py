from __future__ import annotations

import os

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from .models import Match

_COLUMN_LABELS = {
    "league": "League",
    "time": "Kick-off Time",
    "status": "Status",
    "home_team": "Home Team",
    "away_team": "Away Team",
    "odds_1": "Odds 1",
    "odds_x": "Odds X",
    "odds_2": "Odds 2",
    "ou_line": "O/U Line",
    "odds_over": "Odds Over",
    "odds_under": "Odds Under",
    "stats_available": "Stats Available",
    "home_rank": "Home Rank",
    "home_points": "Home Points",
    "home_form": "Home Form",
    "away_rank": "Away Rank",
    "away_points": "Away Points",
    "away_form": "Away Form",
    "h2h_home_wins": "H2H Home Wins",
    "h2h_draws": "H2H Draws",
    "h2h_away_wins": "H2H Away Wins",
    "h2h_since": "H2H Since (Year)",
    "match_url": "Match Link",
}

_ODDS_COLUMNS = {"odds_1", "odds_x", "odds_2", "odds_over", "odds_under"}


def matches_to_dataframe(matches: list[Match]) -> pd.DataFrame:
    return pd.DataFrame(m.to_flat_dict() for m in matches)


def _style_sheet(ws, df: pd.DataFrame, odds_cols: set[str]) -> None:
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")

    for col_idx, col_key in enumerate(df.columns, start=1):
        label = _COLUMN_LABELS.get(col_key, col_key)
        cell = ws.cell(row=1, column=col_idx, value=label)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        letter = get_column_letter(col_idx)
        sample_values = [str(v) if pd.notna(v) else "" for v in df[col_key].head(200).tolist()]
        max_len = max([len(label)] + [len(v) for v in sample_values])
        ws.column_dimensions[letter].width = min(max(max_len + 2, 10), 42)

        if col_key in odds_cols:
            for row_idx in range(2, len(df) + 2):
                ws.cell(row=row_idx, column=col_idx).number_format = "0.00"

    ws.freeze_panes = "A2"

    if len(df) > 0:
        last_col_letter = get_column_letter(len(df.columns))
        table_ref = f"A1:{last_col_letter}{len(df) + 1}"
        table = Table(displayName="MatchesTable", ref=table_ref)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
            showRowStripes=True, showColumnStripes=False,
        )
        ws.add_table(table)


def save_to_excel(matches: list[Match], path: str) -> None:
    parent = os.path.dirname(path)
    if parent:
        os.makedirs(parent, exist_ok=True)

    df = matches_to_dataframe(matches)

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="Matches", index=False, header=False, startrow=1)
        _style_sheet(writer.sheets["Matches"], df, _ODDS_COLUMNS)
